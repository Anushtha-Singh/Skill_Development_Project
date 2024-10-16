from flask import Flask, render_template, request, send_file, url_for
import pdfplumber
import pandas as pd
import matplotlib.pyplot as plt
import os
from io import BytesIO
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image
import tempfile
from docx import Document  # For Word (.docx) files
import xlrd  # For Excel files (.xls)

app = Flask(__name__)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload():
    if 'file' not in request.files:
        return 'No file part', 400

    file = request.files['file']
    if file.filename == '':
        return 'No selected file', 400

    file_ext = os.path.splitext(file.filename)[1].lower()  # Get file extension
    if file_ext == '.pdf':
        return handle_pdf(file)
    elif file_ext == '.docx':
        return handle_word(file)
    elif file_ext in ['.xls', '.xlsx']:
        return handle_excel(file)
    else:
        return 'Unsupported file type', 400

# Function to handle PDF files
def handle_pdf(file):
    pdf_path = tempfile.mktemp(suffix='.pdf')
    file.save(pdf_path)

    tables = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            table = page.extract_table()
            if table:
                df_page = pd.DataFrame(table[1:], columns=table[0])
                tables.append(df_page)

    if tables:
        df = pd.concat(tables, ignore_index=True)
        return generate_charts_and_excel(df)
    else:
        return 'No tables found in the PDF.', 400

# Function to handle Word (.docx) files
def handle_word(file):
    doc_path = tempfile.mktemp(suffix='.docx')
    file.save(doc_path)

    doc = Document(doc_path)
    tables = []
    for table in doc.tables:
        data = []
        for row in table.rows:
            data.append([cell.text.strip() for cell in row.cells])
        if data:
            df = pd.DataFrame(data[1:], columns=data[0])  # First row as header
            tables.append(df)

    if tables:
        df = pd.concat(tables, ignore_index=True)
        return generate_charts_and_excel(df)
    else:
        return 'No tables found in the Word document.', 400

# Function to handle Excel (.xls and .xlsx) files
def handle_excel(file):
    df = pd.read_excel(file)
    return generate_charts_and_excel(df)

# Function to generate charts and save data to Excel
def generate_charts_and_excel(df):
    column_name = request.form['column_name']
    df[column_name] = pd.to_numeric(df[column_name], errors='coerce')
    df = df.dropna(subset=[column_name])

    categories = pd.cut(df[column_name], bins=[0, 30, 40, 50, 60, 70, 80, 90, 100], include_lowest=True)
    counts = categories.value_counts(sort=False)

    pie_chart_path = os.path.join('static', 'pie_chart.png')
    line_chart_path = os.path.join('static', 'line_chart.png')
    bar_chart_path = os.path.join('static', 'bar_chart.png')  # Path for bar chart

    # Pie chart
    plt.figure(figsize=(6, 6))
    plt.pie(counts, labels=counts.index, autopct='%1.1f%%', startangle=140)
    plt.title(f'Pie Chart of {column_name}')
    plt.savefig(pie_chart_path)
    plt.close()

    # Line chart
    plt.figure(figsize=(10, 4))
    plt.plot(counts.index.astype(str), counts.values, marker='o')
    plt.title(f'Line Chart of {column_name}')
    plt.xlabel('Ranges')
    plt.ylabel('Count')
    plt.savefig(line_chart_path)
    plt.close()

    # Bar chart
    plt.figure(figsize=(10, 6))
    plt.bar(counts.index.astype(str), counts.values, color='skyblue')
    plt.title(f'Bar Chart of {column_name}')
    plt.xlabel('Ranges')
    plt.ylabel('Count')
    plt.savefig(bar_chart_path)
    plt.close()

    # Create Excel file
    excel_path = os.path.join('static', 'data.xlsx')
    wb = openpyxl.Workbook()
    ws = wb.active

    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)

    # Add pie chart to Excel
    img_pie = Image(pie_chart_path)
    img_pie.width, img_pie.height = 300, 300
    ws.add_image(img_pie, 'H2')

    # Add line chart to Excel
    img_line = Image(line_chart_path)
    img_line.width, img_line.height = 400, 200
    ws.add_image(img_line, 'H20')

    # Add bar chart to Excel
    img_bar = Image(bar_chart_path)
    img_bar.width, img_bar.height = 400, 200
    ws.add_image(img_bar, 'H35')

    wb.save(excel_path)

    excel_html = df.to_html(classes="table table-striped", index=False)
    return render_template('result.html', excel_html=excel_html,
                           pie_chart=url_for('static', filename='pie_chart.png'),
                           line_chart=url_for('static', filename='line_chart.png'),
                           bar_chart=url_for('static', filename='bar_chart.png'),
                           excel_file='data.xlsx')

@app.route('/download/<filename>')
def download(filename):
    return send_file(os.path.join('static', filename), as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)
